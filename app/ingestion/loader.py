
from pathlib import Path
import csv

import pandas as pd


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".csv",
}


def _validate_file(file_path: str | Path) -> Path:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    if not path.is_file():
        raise ValueError(f"Path is not a file: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    return path


def load_text_file(file_path: str | Path) -> str:
    """Load a UTF-8 text file."""
    path = _validate_file(file_path)
    return path.read_text(encoding="utf-8-sig")


def load_pdf(file_path: str | Path) -> str:
    """Extract PDF text using PyMuPDF, with pypdf as a fallback."""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError("load_pdf requires a PDF file.")

    errors = []

    # First choice: PyMuPDF
    try:
        import fitz

        with fitz.open(path) as document:
            pages = [
                f"[Page {index}]\n{text.strip()}"
                for index, page in enumerate(document, start=1)
                if (text := page.get_text("text").strip())
            ]

        if pages:
            return "\n\n".join(pages)

        errors.append("PyMuPDF extracted no text.")

    except Exception as exc:
        errors.append(f"PyMuPDF: {exc}")

    # Fallback: pypdf
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []

        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()

            if text:
                pages.append(f"[Page {index}]\n{text}")

        if pages:
            return "\n\n".join(pages)

        errors.append("pypdf extracted no text.")

    except Exception as exc:
        errors.append(f"pypdf: {exc}")

    raise ValueError(
        "Could not extract text from PDF. "
        "It may be scanned or image-only. "
        + " | ".join(errors)
    )


def load_docx(file_path: str | Path) -> str:
    """Extract paragraphs and tables from a Word document."""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".docx":
        raise ValueError("load_docx requires a DOCX file.")

    from docx import Document

    document = Document(str(path))
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        rows = []

        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " / ")
                for cell in row.cells
            ]
            rows.append(" | ".join(cells))

        if rows:
            sections.append(
                f"[Table {table_index}]\n" + "\n".join(rows)
            )

    return "\n\n".join(sections)


def load_spreadsheet(file_path: str | Path) -> str:
    """Extract Excel sheets, column headers, and rows."""
    path = _validate_file(file_path)

    if path.suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError("Expected an XLSX or XLS file.")

    sheets = pd.read_excel(
        path,
        sheet_name=None,
        dtype=object,
    )

    sections = []

    for sheet_name, dataframe in sheets.items():
        if dataframe.empty and len(dataframe.columns) == 0:
            continue

        dataframe = dataframe.fillna("")
        lines = [f"[Sheet: {sheet_name}]"]

        if len(dataframe.columns):
            lines.append(
                "Columns: "
                + " | ".join(str(column) for column in dataframe.columns)
            )

        for row in dataframe.itertuples(index=False, name=None):
            values = [str(value) for value in row]
            lines.append(" | ".join(values))

        sections.append("\n".join(lines))

    return "\n\n".join(sections)


def load_csv(file_path: str | Path) -> str:
    """Extract CSV headers and rows."""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".csv":
        raise ValueError("load_csv requires a CSV file.")

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        rows = list(csv.reader(file))

    return "\n".join(
        " | ".join(cell.strip() for cell in row)
        for row in rows
    )


def load_document(file_path: str | Path) -> str:
    """Dispatch a supported file to its appropriate loader."""
    path = _validate_file(file_path)
    extension = path.suffix.lower()

    if extension == ".txt":
        text = load_text_file(path)
    elif extension == ".pdf":
        text = load_pdf(path)
    elif extension == ".docx":
        text = load_docx(path)
    elif extension in {".xlsx", ".xls"}:
        text = load_spreadsheet(path)
    elif extension == ".csv":
        text = load_csv(path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    if not text.strip():
        raise ValueError(
            f"No extractable text found in {path.name}. "
            "Scanned PDFs may require OCR."
        )

    return text
